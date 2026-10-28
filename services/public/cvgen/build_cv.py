#!/usr/bin/env python3
"""Renders the CVs to HTML and PDF with headless Chrome, and the role CVs also to TXT and DOCX.

Usage: python3 build_cv.py
  out/Kalislav-Smirnov-CV{,-SK,-Engineering}.pdf     downloadable CVs (site root)
  out/cv/Kalislav-Smirnov-<Role>.{pdf,txt,docx}       role CVs served from /cv/

DOCX needs python-docx (`uv run --no-project --with python-docx python3 build_cv.py`); without it the
DOCX files are skipped with a warning. Set CHROME to the browser binary when it is not found on its own.
"""
import html
import os
import pathlib
import shutil
import subprocess
import sys

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from cv_data import CONTACTS, EN, ENG, ROLE_CVS, SK, SKILLS_TAGS  # noqa: E402

HERE = pathlib.Path(__file__).parent
OUT = HERE / "out"
MAC_CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
CHROME = (os.environ.get("CHROME") or (MAC_CHROME if os.path.exists(MAC_CHROME) else None)
          or shutil.which("chromium") or shutil.which("chromium-browser") or shutil.which("google-chrome"))

CSS = """
@page { size: A4; margin: 14mm 15mm 14mm; }
* { box-sizing: border-box; }
body { margin: 0; font: 9.6pt/1.42 Inter, "Helvetica Neue", Arial, sans-serif; color: #1f2937; }
h1 { margin: 0; font-size: 21pt; letter-spacing: -0.02em; color: #111827; }
.role { margin: 2px 0 0; font-size: 11.5pt; font-weight: 600; color: #0f766e; }
.meta { margin: 6px 0 0; color: #4b5563; font-size: 9pt; }
.meta a { color: #1d4ed8; text-decoration: none; }
.meta span + span::before { content: "·"; margin: 0 6px; color: #9ca3af; }
h2 { margin: 14px 0 6px; padding-bottom: 3px; border-bottom: 1px solid #d1d5db; font-size: 10.5pt;
     text-transform: uppercase; letter-spacing: 0.06em; color: #111827; }
p { margin: 0 0 5px; }
.job { margin: 0 0 9px; break-inside: avoid; }
.job__head { display: flex; justify-content: space-between; gap: 12px; }
.job__role { font-weight: 700; color: #111827; }
.job__company { color: #0f766e; font-weight: 600; }
.job__period { color: #6b7280; font-style: italic; white-space: nowrap; padding-right: 2px; }
ul { margin: 3px 0 3px; padding-left: 15px; }
li { margin: 1.5px 0; }
.stack { color: #4b5563; font-size: 8.8pt; }
.stack b { color: #374151; }
.degree { font-weight: 700; color: #111827; }
.note { color: #4b5563; }
.skills { display: grid; grid-template-columns: 150px 1fr; gap: 2px 10px; }
.skills dt { font-weight: 600; color: #374151; }
.skills dd { margin: 0; }
.keep { break-inside: avoid; }
"""


def esc(s):
    return html.escape(s, quote=True)


def render(d):
    contacts = "".join(f'<span><a href="{esc(u)}">{esc(t)}</a></span>' for t, u in CONTACTS)
    jobs = []
    for j in d["jobs"]:
        company = esc(j["company"]) + (f' · {esc(j["type"])}' if j["type"] else "")
        intro = f'<p>{esc(j["intro"])}</p>' if j["intro"] else ""
        points = "".join(f"<li>{esc(p)}</li>" for p in j["points"])
        jobs.append(
            f'<div class="job"><div class="job__head"><div><span class="job__role">{esc(j["role"])}</span>, '
            f'<span class="job__company">{company}</span></div><span class="job__period">{esc(j["period"])}</span></div>'
            f'{intro}<ul>{points}</ul><p class="stack"><b>{esc(d["stack_label"])}:</b> {esc(j["stack"])}</p></div>'
        )
    tags = sorted(SKILLS_TAGS.items(), key=lambda kv: kv[0] != "eng") if d.get("edu_first") else SKILLS_TAGS.items()
    skills = "".join(f'<dt>{esc(d["skills"][k])}</dt><dd>{esc(v)}</dd>' for k, v in tags)
    aero = "".join(f"<li>{esc(p)}</li>" for p in d["aero_points"])
    lab = "".join(f"<li>{esc(p)}</li>" for p in d["lab"])
    summary = "".join(f"<p>{esc(p)}</p>" for p in d["summary"])
    exp = f"<h2>{esc(d['h']['exp'])}</h2>{''.join(jobs)}"
    agents = ""
    if d.get("agents"):
        items = "".join(f'<li><b>{esc(t)}.</b> {esc(body)}</li>' for t, body in d["agents"])
        agents = f'<h2>{esc(d["h"]["agents"])}</h2><ul>{items}</ul>'
    edu = (f'<div class="keep"><h2>{esc(d["h"]["aero"])}</h2>'
           f'<p><span class="degree">{esc(d["degree"])}</span><br><span class="note">{esc(d["degree_note"])}</span></p>'
           f'<p>{esc(d["field"])}</p><ul>{aero}</ul><p class="stack"><b>{esc(d["aero_tools"])}</b></p></div>')
    if d.get("agents_first"):
        blocks = agents + exp + edu
    elif d.get("edu_first"):
        blocks = edu + exp + agents
    else:
        blocks = exp + agents + edu
    return f"""<!doctype html>
<html lang="{d['lang']}"><head><meta charset="utf-8"><title>{esc(d['title'])}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700&display=swap">
<style>{CSS}</style></head><body>
<h1>Kalislav Smirnov</h1>
<p class="role">{esc(d['role'])}</p>
<p class="meta"><span>{esc(d['location'])}</span>{contacts}</p>
<h2>{esc(d['h']['summary'])}</h2>{summary}
{blocks}
<div class="keep"><h2>{esc(d['h']['skills'])}</h2><dl class="skills">{skills}</dl></div>
<div class="keep"><h2>{esc(d['h']['lab'])}</h2><ul>{lab}</ul></div>
<div class="keep"><h2>{esc(d['h']['lang'])}</h2><p>{esc(d['languages'])}</p></div>
</body></html>"""


def outline(d):
    """The CV as (heading, [(kind, text)]) blocks for the plain renderers; kind is sub, p or li."""
    exp = []
    for j in d["jobs"]:
        exp.append(("sub", f'{j["role"]} | {j["company"]}' + (f' | {j["type"]}' if j["type"] else "") + f' | {j["period"]}'))
        exp += [("p", j["intro"])] if j["intro"] else []
        exp += [("li", p) for p in j["points"]]
        exp.append(("p", f'{d["stack_label"]}: {j["stack"]}'))
    edu = [("sub", d["degree"]), ("p", d["degree_note"]), ("p", d["field"]),
           *[("li", p) for p in d["aero_points"]], ("p", d["aero_tools"])]
    agents = [("li", f"{t}. {body}") for t, body in d.get("agents", [])]
    h = d["h"]
    middle = [(h["exp"], exp)] + ([(h["agents"], agents)] if agents else []) + [(h["aero"], edu)]
    if d.get("edu_first"):
        middle = [middle[-1], *middle[:-1]]
    tags = sorted(SKILLS_TAGS.items(), key=lambda kv: kv[0] != "eng") if d.get("edu_first") else SKILLS_TAGS.items()
    return [(h["summary"], [("p", p) for p in d["summary"]]), *middle,
            (h["skills"], [("p", f'{d["skills"][k]}: {v}') for k, v in tags]),
            (h["lab"], [("li", p) for p in d["lab"]]), (h["lang"], [("p", d["languages"])])]


def contact_line():
    return " | ".join(t if u.startswith("mailto:") or t.startswith("Telegram") else u.split("?")[0].rstrip("/")
                      for t, u in CONTACTS)


def render_txt(d):
    lines = ["Kalislav Smirnov", "", d["role"], "", d["location"], contact_line()]
    for heading, items in outline(d):
        lines += ["", heading]
        for kind, text in items:
            lines += ["", f"- {text}" if kind == "li" else text]
    return "\n".join(lines) + "\n"


def render_docx(d, path):
    try:
        import docx
        from docx.shared import Cm, Pt
    except ImportError:
        print(f"skip {path.name}: python-docx is not installed", file=sys.stderr)
        return
    doc = docx.Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Kalislav Smirnov", level=0)
    doc.add_paragraph().add_run(d["role"]).bold = True
    doc.add_paragraph(f'{d["location"]}\n{contact_line()}')
    for heading, items in outline(d):
        doc.add_heading(heading, level=1)
        for kind, text in items:
            if kind == "li":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph().add_run(text).bold = kind == "sub"
    doc.save(path)


def pdf(d, folder, name):
    if not CHROME:
        sys.exit("Chrome/Chromium not found: set CHROME")
    src = folder / f"{name}.html"
    src.write_text(render(d))
    out = folder / f"{name}.pdf"
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
                    f"--user-data-dir={OUT / '.chrome'}", "--virtual-time-budget=4000",
                    f"--print-to-pdf={out}", src.as_uri()],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("wrote", out)


def main():
    OUT.mkdir(exist_ok=True)
    for d, name in [(EN, "Kalislav-Smirnov-CV"), (SK, "Kalislav-Smirnov-CV-SK"), (ENG, "Kalislav-Smirnov-CV-Engineering")]:
        pdf(d, OUT, name)
    roles = OUT / "cv"
    roles.mkdir(exist_ok=True)
    for name, d in ROLE_CVS.items():
        pdf(d, roles, name)
        (roles / f"{name}.txt").write_text(render_txt(d))
        render_docx(d, roles / f"{name}.docx")
        (roles / f"{name}.html").unlink()


if __name__ == "__main__":
    main()
